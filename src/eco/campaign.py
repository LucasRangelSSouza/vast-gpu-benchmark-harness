# Campaign driver: picks offers, runs preflight checks, launches a session.
# Original harness script, copied unmodified apart from this header. Comments inside may be in Portuguese.
"""Uma GPU por vez, com checagens baratas ANTES de ligar qualquer maquina paga.
uso: python -m eco.campaign <GPU> [--airllm] [--items "a b c"]   (roda eco.session_script com remote/run_all.sh)
Checagens locais (custo zero): acesso HF dos repos gated, saldo Vast vs previsto, instancias orfas."""
import json
import os
import subprocess
import sys

sys.path.insert(0, ".")
import httpx  # noqa: E402

from eco import vast  # noqa: E402
from eco.common import USD_BRL  # noqa: E402

CFG = {  # gpu: (ttl_min, US$/h previsto, min_disk)
    "RTX 3090": (150, 0.22, 300), "RTX 4090": (150, 0.43, 300), "RTX 5090": (150, 0.55, 300),
    "L40S": (160, 0.80, 300), "RTX PRO 6000": (170, 1.36, 300), "H100 NVL": (110, 3.03, 300),
}
GATED = {"flux": "black-forest-labs/flux-3-image-klein", "personaplex": "nvidia/personaplex-7b-v1", "supra": "facebook/dinov3-vitl16-pretrain-lvd1689m"}
ALL = "libs q9 awq27 fp8_plain fp8_dflash bf16_dflash nvfp4_dflash flux personaplex supra".split()


ROWMAP = {"NVFP4": "nvfp4_dflash", "FP8 + DFlash2": "fp8_dflash", "FP8 (sem DFlash2)": "fp8_plain", "BF16 + DFlash2": "bf16_dflash", "AWQ": "awq27", "Qwen3.5-9B": "q9",
          "YOLO": "yolo2", "SAM 2": "sam2", "Grounding DINO": "gdino", "Mask R-CNN": "maskrcnn_tv", "SDXL": "sdxl", "LTX": "ltx", "F5-TTS": "whisper_f5", "FFmpeg": "ffmpeg",
          "AirLLM": "airllm", "FLUX": "flux", "PersonaPlex": "personaplex", "SUPRA": "supra"}
LIBNAMES = {"yolo2", "sam2", "gdino", "maskrcnn_tv", "sdxl", "ltx", "whisper_f5", "ffmpeg"}


def missing_items(gpu, items):
    """So entra no manifest o que ainda e ESTIMATIVA na matriz desta GPU (o medido nao e repetido)."""
    import subprocess as sp
    sp.run([sys.executable, "eco/aggregate.py"], capture_output=True)
    sp.run([sys.executable, "eco/compat.py"], capture_output=True)
    import openpyxl
    ws = openpyxl.load_workbook("data/derived/COMPAT_MATRIX.xlsx").active
    hdr = [c.value for c in ws[1]]
    col = next(i for i, h in enumerate(hdr) if h and h.startswith(gpu))
    need = set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[col] and str(r[col]).startswith("ESTIMATIVA"):
            for k, v in ROWMAP.items():
                if k in str(r[1]):
                    need.add(v)
    libs = [x for x in need if x in LIBNAMES]
    out = [i for i in items if i in need or (i == "libs" and libs)]
    print(f"medido e pulado: {sorted(set(items) - set(out) - ({'libs'} if libs else set()))}; a medir: {out}; libs: {libs}")
    return out, libs


def hf_ok(repo, tok):
    r = httpx.head(f"https://huggingface.co/{repo}/resolve/main/.gitattributes", headers={"Authorization": "Bearer " + tok}, timeout=30)
    return r.status_code in (200, 302, 307)


def main():
    gpu = sys.argv[1]
    ttl, price, disk = CFG[gpu]
    ttl = int(os.environ.get("TTL_MIN", ttl))
    items = ALL[:]
    if "--items" in sys.argv:
        items = sys.argv[sys.argv.index("--items") + 1].split()
    extras = "--extras" in sys.argv
    entry_arg = sys.argv[sys.argv.index("--entry") + 1] if "--entry" in sys.argv else None
    if entry_arg:
        items, libs = [], []
    elif extras:
        items, libs = ["supra"], []
    else:
        items, libs = missing_items(gpu, items)
    tok = json.load(open(os.path.join(os.environ.get("SECRETS_DIR", "secrets"), "huggingface.json")))["HF_TOKEN"]
    for it, repo in GATED.items():
        if it in items and not hf_ok(repo, tok):
            print(f"ACESSO NEGADO {repo}: item {it} removido (aceitar a licenca em huggingface.co/{repo})")
            items.remove(it)
    credit = vast.user()["credit"]
    est = price * ttl * 0.85 / 60
    print(f"saldo US$ {credit:.2f} (R$ {credit * USD_BRL:.0f}); previsto p/ {gpu}: US$ {est:.2f} (R$ {est * USD_BRL:.0f}); itens: {items}")
    orphans = [x for x in vast.instances() if x.get("actual_status") != "exited"]
    if orphans:
        print("INSTANCIAS ATIVAS, abortando:", [(x["id"], x.get("label")) for x in orphans])
        sys.exit(2)
    if credit < est + 1.0:
        print("saldo insuficiente para esta GPU"); sys.exit(3)
    GB = {"libs": 25, "supra": 8, "q9": 18, "awq27": 19, "fp8_plain": 52, "fp8_dflash": 4, "bf16_dflash": 0, "nvfp4_dflash": 33, "flux": 60, "personaplex": 18, "supra": 8}
    dl_gb = sum(GB.get(i, 0) for i in items) + 10 + (120 if (extras or entry_arg) else 0)
    print(f"download previsto ~{dl_gb} GB")
    env = dict(os.environ, DL_GB=str(dl_gb), PASS_HF="1", DISK_GB="330", MIN_DISK=str(disk), BUDGET_BRL=os.environ.get("BUDGET_BRL", "9999"), SOFT_STOP_BRL=os.environ.get("SOFT_STOP_BRL", "9998"))
    pre = ("LIBS=\"%s\"; export LIBS; " % " ".join(libs)) + "ITEMS=\"%s\"; export ITEMS; " % " ".join(items) + ("AIRLLM=1; export AIRLLM; " if "--airllm" in sys.argv else "")
    entry = entry_arg or ("run_extras.sh" if extras else "run_all.sh")
    open("remote/_all_launch.sh", "w", newline="\n").write("#!/bin/bash\n" + pre + "bash /workspace/eco/" + entry + "\n")
    if os.environ.get("DRY"):
        print("DRY: launcher gerado:", open("remote/_all_launch.sh").read()); sys.exit(0)
    tag = "all-" + gpu.replace(" ", "")
    sys.exit(subprocess.call([sys.executable, "-m", "eco.session_script", gpu, tag, "_all_launch.sh", str(ttl), "vastai/vllm:v0.30.0-cuda-13.0", "13.0"], env=env))


main()
