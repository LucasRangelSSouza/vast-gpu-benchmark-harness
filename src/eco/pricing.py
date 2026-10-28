# Cost-per-unit inputs for pricing, from measured data.
# Original harness script, copied unmodified apart from this header. Comments inside may be in Portuguese.
"""Insumos de precificacao do modelo 'VM sem GPU + GPU sob demanda repassada ao cliente'.
Custo por unidade de trabalho = (US$/h da GPU) x (segundos de GPU) / 3600, MAIS o custo fixo de ligar a GPU (boot + baixar modelo) por sessao.
Tudo a partir do medido (libs_summary, llm_summary, leases). Saida: data/derived/PRICING_INPUTS.xlsx e .md"""
import csv
import glob
import json
import statistics as st
import sys

sys.path.insert(0, ".")
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from eco.common import DATA, RAW, USD_BRL  # noqa: E402

DER = DATA / "derived"
libs = [json.loads(line) for line in open(DER / "libs_summary.jsonl", encoding="utf-8")]
llm = [json.loads(line) for line in open(DER / "llm_summary.jsonl", encoding="utf-8")]

PRICE = {"RTX 4090": 0.43, "RTX 5090": 0.55, "L40S": 0.80, "RTX PRO 6000": 1.36, "H100": 2.80, "RTX 3090": 0.22, "GB10": 0.449}


def gname(s):
    s = s or ""
    for k in PRICE:
        if k.replace("RTX ", "") in s or k in s:
            return k
    return s.split(",")[0].replace("NVIDIA ", "").replace("GeForce ", "").strip()


def price(g):
    return PRICE.get(g)


# ---- boot: criar -> ssh pronto, por GPU (leases)
boot = {}
for d in glob.glob(str(RAW / "leases" / "2026*")):
    try:
        ev = [json.loads(line) for line in open(d + "/lease_events.jsonl", encoding="utf-8")]
        man = json.load(open(d + "/manifest.json", encoding="utf-8"))
    except Exception:  # noqa: BLE001
        continue
    cr = next((e for e in ev if e["event"] == "create_response" and isinstance(e.get("resp"), dict) and e["resp"].get("success")), None)
    rd = next((e for e in ev if e["event"] == "ssh_ready"), None)
    if cr and rd:
        from datetime import datetime
        s = (datetime.fromisoformat(rd["t_utc"]) - datetime.fromisoformat(cr["t_utc"])).total_seconds()
        g = gname((man.get("offer_snapshot") or {}).get("gpu_name"))
        boot.setdefault(g, []).append(s)
fails = {}
for d in glob.glob(str(RAW / "leases" / "2026*")):
    try:
        ev = [json.loads(line) for line in open(d + "/lease_events.jsonl", encoding="utf-8")]
        man = json.load(open(d + "/manifest.json", encoding="utf-8"))
    except Exception:  # noqa: BLE001
        continue
    g = gname((man.get("offer_snapshot") or {}).get("gpu_name"))
    t = fails.setdefault(g, [0, 0])
    if any(e["event"] == "create_response" for e in ev):
        t[1] += 1
        if any(e["event"] in ("not_running_timeout", "ssh_timeout") for e in ev):
            t[0] += 1

rows = []


def add(cap, unit, gpu, usd_per_unit, basis, note=""):
    if usd_per_unit is None:
        return
    rows.append({"capacidade": cap, "unidade": unit, "gpu": gpu, "usd_gpu_h": price(gpu), "usd_por_unidade": usd_per_unit, "brl_por_unidade": usd_per_unit * USD_BRL, "base_medida": basis, "obs": note})


# ---- imagem (SDXL-Turbo via diffusers, 512x512, 4 passos, lote 4)
for r in libs:
    g = gname(r.get("gpu"))
    if r["lib"] == "sdxl_turbo" and r["status"] == "ok" and r.get("batch") == 4 and price(g):
        add("Imagem 512px (SDXL-Turbo)", "1000 imagens", g, price(g) * r["s_per_image"] / 3600 * 1000, f"{r['s_per_image']} s/imagem")
    if r["lib"] == "ltx_video" and r["status"] == "ok" and r.get("frames") == 49 and price(g):
        add("Video gerado (LTX-Video 512x320)", "1 min de video", g, price(g) * r["gen_s_per_video_s"] * 60 / 3600, f"{r['gen_s_per_video_s']} s de GPU por s de video")
    if r["lib"] == "f5_tts" and r["status"] == "ok" and r.get("case") == "medio" and r["rtf"] < 1 and price(g):
        add("Voz sintetica com clonagem (F5-TTS)", "1 hora de audio", g, price(g) * r["rtf"], f"RTF {r['rtf']}")
    if r["lib"] == "yolo_gpu_only" and r["status"] == "ok" and r.get("model") == "yolo11m" and r.get("batch") == 16 and r.get("imgsz") == 640 and price(g):
        add("Deteccao de objetos (YOLO11m, so GPU)", "1 milhao de frames", g, price(g) * 1e6 / r["fps_gpu"] / 3600, f"{r['fps_gpu']} fps")
    if r["lib"] == "sam2" and r["status"] == "ok" and price(g):
        add("Segmentacao (SAM 2 hiera_large)", "1000 imagens", g, price(g) * r["p50_ms"] / 1000 * 1000 / 3600, f"{r['p50_ms']} ms/imagem")
    if r["lib"] == "grounding_dino" and r["status"] == "ok" and price(g):
        add("Deteccao por texto (Grounding DINO)", "1000 imagens", g, price(g) * r["p50_ms"] / 1000 * 1000 / 3600, f"{r['p50_ms']} ms/imagem")
    if r["lib"] == "ffmpeg" and r["status"] == "ok" and r.get("codec") == "hevc_nvenc" and r.get("res") == "4k" and price(g):
        add("Transcodificacao 4K HEVC (NVENC)", "1 hora de video", g, price(g) / r["speed_x"], f"{r['speed_x']}x tempo real")
    if r["lib"] == "maskrcnn_torchvision" and r["status"] == "ok" and price(g):
        add("Mask R-CNN (torchvision, proxy Detectron2)", "1000 imagens", g, price(g) * r["p50_ms"] / 1000 * 1000 / 3600, f"{r['p50_ms']} ms/imagem")
    if r["lib"] in ("faster_whisper", "faster_whisper_gpu") and r["status"] == "ok" and r.get("rtf") and price(g):
        add("Transcricao de audio (faster-whisper)", "1 hora de audio", g, price(g) * r["rtf"], f"RTF {r['rtf']}")

# ---- LLM (tokens): por modelo/GPU, no nivel de maior vazao dentro da meta
best = {}
for r in llm:
    if r["thinking"] or r["errors"] > r["requests"] * 0.4:
        continue
    g = gname(r.get("gpu"))
    k = (r["model"], g, r["variant"])
    if r["agg_out_tps"] and (k not in best or r["agg_out_tps"] > best[k]["agg_out_tps"]):
        best[k] = r
for (model, g, var), r in best.items():
    p = price(g) or r.get("dph_usd")
    if not p:
        continue
    add(f"LLM {model} ({var})", "1 milhao de tokens de saida", g, p / (r["agg_out_tps"] * 3600) * 1e6, f"{r['agg_out_tps']} tok/s agregado com {r['users']} usuarios", "GPU saturada; ocupacao baixa encarece proporcionalmente")

rows.sort(key=lambda x: (x["capacidade"], x["usd_por_unidade"]))
md = ["# PRICING_INPUTS: custo de GPU sob demanda por tarefa (medido) — NAO e preco de venda\n",
      f"Modelo de negocio: o Lunar vende a VM sem GPU; a GPU e alugada na hora, o Lunar paga e o cliente reembolsa. USD_BRL={USD_BRL}. Preco da GPU = US$/h do marketplace no dia da medicao (varia).",
      "`custo por unidade = US$/h x segundos de GPU / 3600`. **Some o custo de ligar a GPU** (tabela de boot abaixo) em cada sessao. Margem e IOF/cambio ficam por sua conta.\n",
      "## Custo fixo de ligar a GPU (criar -> SSH pronto, medido)\n", "| GPU | sessoes | boot p50 (s) | boot max (s) | falhas de subir/sessoes |", "|---|---:|---:|---:|---:|"]
for g, v in sorted(boot.items(), key=lambda kv: st.median(kv[1])):
    f = fails.get(g, [0, 0])
    md.append(f"| {g} | {len(v)} | {st.median(v):.0f} | {max(v):.0f} | {f[0]}/{f[1]} |")
md += ["", "O boot acima NAO inclui baixar o modelo: o 27B em 4 bits (18 GB) levou ~4-8 min e o BF16 (52 GB) ~7-12 min em hosts de ~600 Mbps; em host lento chegou a 17 GB em 65 min (cancelado). Use volume persistente/cache quando o mesmo cliente repete a tarefa.", "",
       "## Custo por unidade de trabalho (menor primeiro, por capacidade)\n", "| capacidade | unidade | GPU | US$/h | US$/unidade | R$/unidade | base medida |", "|---|---|---|---:|---:|---:|---|"]
last = None
for r in rows:
    md.append(f"| {r['capacidade']} | {r['unidade']} | {r['gpu']} | {r['usd_gpu_h']} | {r['usd_por_unidade']:.4f} | {r['brl_por_unidade']:.3f} | {r['base_medida']} |")
(DER / "PRICING_INPUTS.md").write_text("\n".join(md) + "\n", encoding="utf-8")

wb = Workbook()
ws = wb.active
ws.title = "Custo por unidade"
ws.append(["Capacidade", "Unidade", "GPU", "US$/h GPU", "US$/unidade", "R$/unidade", "Base medida", "Obs"])
for r in rows:
    ws.append([r["capacidade"], r["unidade"], r["gpu"], r["usd_gpu_h"], round(r["usd_por_unidade"], 5), round(r["brl_por_unidade"], 4), r["base_medida"], r["obs"]])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="305496")
b = wb.create_sheet("Boot da GPU")
b.append(["GPU", "Sessoes", "Boot p50 (s)", "Boot max (s)", "Falhas de subir", "Sessoes criadas"])
for g, v in sorted(boot.items()):
    f = fails.get(g, [0, 0])
    b.append([g, len(v), round(st.median(v)), round(max(v)), f[0], f[1]])
for sh, widths in ((ws, (44, 28, 14, 10, 12, 12, 46, 50)), (b, (16, 10, 14, 14, 16, 16))):
    for i, w in enumerate(widths):
        sh.column_dimensions[chr(65 + i)].width = w
wb.save(DER / "PRICING_INPUTS.xlsx")
print(len(rows), "linhas de custo;", {g: len(v) for g, v in boot.items()})
